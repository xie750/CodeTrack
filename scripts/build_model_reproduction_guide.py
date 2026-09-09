from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MATERIAL_ROOT = ROOT / "项目申报材料" / "05-作品代码"
OUT_PATH = MATERIAL_ROOT / "05_复现说明" / "CodeTrack_模型复现说明.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "000000", size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def remove_style_borders(style) -> None:
    p_pr = style._element.pPr
    if p_pr is None:
        return
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF", size=9)
        set_cell_shading(table.rows[0].cells[index], "1F4E79")
        set_cell_borders(table.rows[0].cells[index])
        if widths:
            table.rows[0].cells[index].width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, size=9)
            set_cell_shading(cells[index], "F3F6F9" if row_index % 2 else "FFFFFF")
            set_cell_borders(cells[index])
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(text)


def add_number(doc: Document, number: int, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(f"{number}. {text}")


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def build() -> Path:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        remove_style_borders(style)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("CodeTrack 模型复现说明")
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("面向作品代码材料提交与技术验收")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(89, 89, 89)

    intro = doc.add_paragraph()
    intro.add_run("说明：").bold = True
    intro.add_run(
        "本说明依据当前 CodeTrack 项目源码、部署配置、模型接入代码和已提供的 GGUF 文件整理。"
        "现阶段可以复现应用层的模型接入、课程知识库增强、受控输出和规则兜底链路；"
        "由于未随材料提供模型训练数据和训练过程记录，不能据此复现模型训练过程。"
    )

    doc.add_heading("一 材料现状", level=1)
    add_table(
        doc,
        ["材料项", "当前情况", "是否可直接复现"],
        [
            ["模型文件", "已提供 codetrack-q4_k_m.gguf，文件头为 GGUF", "可以归档并供模型服务加载"],
            ["作品源码", "工作区包含 backend、teacher_backend、frontend、sandbox、tests 和 deploy", "可以复现应用工程"],
            ["模型 ServiceID", "当前未提供在线服务地址或 ServiceID", "暂不能直接进行在线模型验收"],
            ["微调数据", "当前未发现数据集、字段说明、清洗脚本或授权说明", "不能复现训练数据处理"],
            ["训练过程", "当前未发现训练脚本、超参数、日志、中间检查点或评测集", "不能复现模型训练过程"],
        ],
        [1.35, 3.45, 1.55],
    )

    doc.add_heading("二 模型文件信息", level=1)
    add_table(
        doc,
        ["字段", "内容"],
        [
            ["文件名", "codetrack-q4_k_m.gguf"],
            ["类型", "GGUF 本地模型文件"],
            ["大小", "1,929,902,528 字节，约 1.80 GiB"],
            ["SHA-256", "488E0503BEAA52A591CB400838A74CA111DC35F4FDC7466531236779B2E5AB18"],
            ["归档目录", "02_模型文件或模型ServiceID"],
        ],
        [1.35, 5.0],
    )
    p = doc.add_paragraph()
    p.add_run("文件命名说明：").bold = True
    p.add_run(
        "文件名中的 q4_k_m 表示采用了 4-bit K-M 量化命名约定。仅凭文件名不能证明模型的训练来源、"
        "微调数据或最终效果，相关结论需要后续补充训练和评测材料。"
    )

    doc.add_heading("三 应用层复现链路", level=1)
    p = doc.add_paragraph()
    p.add_run("CodeTrack 的模型调用位于一个受控的应用链路中：").bold = True
    p.add_run("学生提交代码后，系统先获得沙箱编译和测试事实，再将任务上下文、失败测试证据和课程知识源交给模型层。")
    add_code(doc, "学生提交 -> 沙箱编译与测试 -> 构造诊断 payload -> 模型网关或兼容接口 -> 输出校验 -> AI 诊断或规则兜底")
    add_bullet(doc, "任务上下文包括课程、任务、语言、接口规范和学习目标。")
    add_bullet(doc, "工具证据包括测试用例、期望结果摘要、实际输出和错误标签。")
    add_bullet(doc, "知识源引用必须来自允许的课程知识源，不能由模型自行编造 ID。")
    add_bullet(doc, "输出需经过 Schema、引用、置信度和提示泄漏校验；失败时回退到规则诊断。")

    doc.add_heading("四 模型接入方式", level=1)
    add_number(doc, 1, "准备模型服务。使用兼容 GGUF 的推理服务加载 codetrack-q4_k_m.gguf，并将模型文件挂载到服务容器的 /models 目录。")
    add_number(doc, 2, "对外提供 OpenAI 兼容的 /v1/chat/completions 接口，或实现 CodeTrack 内部诊断网关约定的 JSON 接口。")
    add_number(doc, 3, "在 CodeTrack 后端配置模型服务地址、模型名称和访问密钥。密钥只放在部署环境变量中，不写入提交材料。")
    add_number(doc, 4, "初始化演示数据并执行登录、任务提交、测试、诊断和知识源引用验证。")

    p = doc.add_paragraph()
    p.add_run("配置入口：").bold = True
    p.add_run(".env.example、backend/app/core/config.py、backend/app/services/model_gateway.py 和 deploy/docker-compose.yml。")
    add_code(doc, "CODETRACK_MODEL_GATEWAY_URL=<内部诊断网关地址>\nCODETRACK_MODEL_API_BASE_URL=<OpenAI兼容服务地址>\nCODETRACK_MODEL_NAME=<服务端模型名称>\nCODETRACK_FINE_TUNED_MODEL_NAME=/models/codetrack-q4_k_m.gguf")
    p = doc.add_paragraph()
    p.add_run("注意：").bold = True
    p.add_run("上述配置项是接入协议，不是当前已经部署的服务地址。当前材料没有可填写的模型 ServiceID。")

    doc.add_heading("五 当前可执行的复现步骤", level=1)
    add_number(doc, 1, "进入作品源码目录 D:\\shy\\CodeTrack。")
    add_number(doc, 2, "准备 Python 和 Node.js 运行环境，安装项目依赖。")
    add_number(doc, 3, "复制 .env.example 为部署环境配置文件；模型服务未准备好时，模型配置可以暂留空。")
    add_number(doc, 4, "执行 python scripts/seed_demo.py，准备课程、任务、知识点、学生学习记录和教师端演示数据。")
    add_number(doc, 5, "启动后端和前端，在学生端使用王同学演示已有学习数据；在教师端使用王老师演示已有教学数据。")
    add_number(doc, 6, "模型服务可用后，使用失败代码提交触发诊断，检查模型输出中的错因、置信度、测试证据和课程知识源引用。")
    add_number(doc, 7, "模型服务不可用时，检查规则兜底结果是否仍引用真实测试结果和课程知识源，并标记为 RULE_FALLBACK。")

    doc.add_heading("六 验收关注点", level=1)
    add_table(
        doc,
        ["验收环节", "应看到的结果"],
        [
            ["模型文件", "文件可被模型服务读取，SHA-256 与材料记录一致"],
            ["接口调用", "模型服务返回 JSON，CodeTrack 后端可以接收并解析"],
            ["输出约束", "诊断类型、置信度、证据引用和知识源引用通过校验"],
            ["知识增强", "回答或诊断能引用当前课程允许的知识源"],
            ["失败降级", "模型不可用或输出不合格时，使用规则兜底且不覆盖测试事实"],
            ["应用闭环", "提交、诊断、提示、总结、画像、资料和推荐链路可继续运行"],
        ],
        [1.55, 4.8],
    )

    doc.add_heading("七 待补充材料", level=1)
    p = doc.add_paragraph()
    p.add_run("为完成模型训练过程复现和完整模型验收，后续建议补充：").bold = True
    for item in [
        "微调数据集或可核验的数据清单，包括来源、授权范围、脱敏方式和版本号。",
        "数据清洗、格式转换、训练集与验证集划分脚本。",
        "基础模型名称和版本、训练框架、量化方式、训练超参数、硬件环境和训练日志。",
        "模型服务启动命令、镜像或依赖清单、健康检查地址和在线 ServiceID。",
        "独立评测集、评价指标、基线模型和对比实验结果。",
        "至少一组可脱敏的输入、模型原始输出、校验结果和最终页面展示结果。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("八 结论", level=1)
    doc.add_paragraph(
        "当前材料足以说明 CodeTrack 的应用层模型接入方案，并能够归档一个可供模型服务加载的 GGUF 文件。"
        "项目可以在没有在线模型服务时通过规则兜底完成基础演示；接入模型服务后，可进一步验收模型诊断和课程知识库引用。"
        "模型训练过程、微调数据和在线服务地址不在当前已提供材料中，待补齐后再更新本说明。"
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("CodeTrack 作品代码材料")
    footer_run.font.name = "Microsoft YaHei"
    footer_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(128, 128, 128)

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
