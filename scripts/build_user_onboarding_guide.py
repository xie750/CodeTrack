from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "docs" / "CodeTrack_学生端与教师端快速上手引导.docx"
SHOT_DIR = ROOT / "docs" / "evidence" / "user-guide"


BLUE = "176CF5"
LIGHT_BLUE = "EAF3FF"
PALE = "F7FAFF"
GRAY_BORDER = "D9D9D9"
TEXT = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRAY_BORDER) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_style_borders(style) -> None:
    p_pr = style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_para(doc: Document, text: str = "", style: str | None = None, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.18
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 10.5, True, TEXT)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, 10.5, False, TEXT)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 10.5, False, TEXT)
    return paragraph


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        set_run_font(run, 16 if level == 1 else 12.5, True, RGBColor(0, 0, 0))
    return paragraph


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, cell in enumerate(table.rows[0].cells):
        cell.text = headers[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, 9.2, True, TEXT)
        cell.width = Inches(widths[idx])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx], 130, 120, 130, 120)
            if row_index % 2 == 1:
                set_cell_shading(cells[idx], PALE)
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cells[idx].width = Inches(widths[idx])
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.line_spacing = 1.12
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, 8.8, False, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent = Cm(0.55)
        for run in paragraph.runs:
            run.text = ""
        run = paragraph.add_run(item)
        set_run_font(run, 10.2, False, TEXT)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent = Cm(0)
        run = paragraph.add_run(f"{index}.  {item}")
        set_run_font(run, 10.2, False, TEXT)


def add_picture(doc: Document, filename: str, caption: str) -> None:
    image_path = SHOT_DIR / filename
    if not image_path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.55))
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(10)
    run = caption_p.add_run(caption)
    set_run_font(run, 8.8, False, MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    remove_style_borders(title)

    for name in ("Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)


def build() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_borders(title)
    run = title.add_run("CodeTrack 学生端与教师端快速上手引导")
    set_run_font(run, 22, True, RGBColor(0, 0, 0))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("面向人工智能专业助学平台的用户入口说明与核心流程指南")
    set_run_font(run, 11, False, MUTED)

    add_para(
        doc,
        "本文档面向第一次接触 CodeTrack 的学生、教师、评审和项目申报材料阅读者。它的目标不是解释所有后台实现，而是把平台中最容易迷路的入口、页面关系和高频操作串成可直接照着走的路线。",
    )
    add_para(
        doc,
        "CodeTrack 当前定位为面向人工智能专业建设的学科垂直大模型与创新助学应用。学生端围绕课程任务、自主学习、AI 助学、学习资料沉淀和学习画像更新形成闭环；教师端围绕课程建设、任务发布、智能批改、学情诊断和教学改进提供支撑。",
    )
    add_para(
        doc,
        "第三项项目材料需要提供 Demo 体验地址和演示账号。本项目将账号分成“演示数据账号”和“初始状态账号”两类，方便评委既能看到完整闭环，也能看到首次登录时的空状态。",
    )

    add_heading(doc, "一 先从共同入口开始", 1)
    add_para(doc, "打开平台后先进入登录页。演示环境可直接使用登录页下方的演示账号，登录后系统会根据角色自动进入学生端或教师端。")
    add_table(
        doc,
        ["项目", "说明"],
        [
            ["Demo 体验地址", "https://demo.codetrack.example.com/  （待部署后替换为正式地址）"],
            ["本地开发地址", "http://127.0.0.1:5173/  （仅用于本机调试，不作为最终提交地址）"],
            ["统一默认密码", "codetrack123"],
        ],
        [1.45, 4.95],
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_para(
        doc,
        "评委查看时建议先使用有数据账号快速理解平台能力，再切换到初始状态账号查看新用户首次进入平台时看到的页面状态。",
    )
    add_table(
        doc,
        ["角色", "账号", "用户", "账号状态", "评委查看重点"],
        [
            ["学生", "wang", "王同学", "演示数据", "查看课程任务、提交记录、AI 诊断、学习画像、资料沉淀和推荐路径"],
            ["学生", "liu", "刘同学", "初始状态", "查看首次登录后的空状态、入口导航和后续加入课程后的使用起点"],
            ["教师", "teacher_wang", "王老师", "演示数据", "查看课程、班级、任务发布、批改、资料管理、知识图谱和学情分析"],
            ["教师", "teacher_li", "李老师", "初始状态", "查看教师首次进入平台时的空课程、空任务和初始化工作台状态"],
        ],
        [0.75, 1.25, 0.95, 1.15, 2.3],
    )
    add_para(
        doc,
        "账号状态定义：王同学和王老师用于展示已沉淀数据的完整演示链路；刘同学仅保留基础课程入口、班级归属和一个未开始任务，不预置提交、画像、AI 会话、资料、推荐或科研实践数据；李老师仅保留可登录账号，不预置教师课程、班级、任务、资料或学情分析数据。",
    )
    add_picture(doc, "common-00-login.png", "图 1 统一登录入口和演示账号位置")

    add_heading(doc, "二 学生端入口地图", 1)
    add_para(
        doc,
        "学生端第一屏不是复杂后台，而是三个服务入口。首次使用时优先选择“我的课程”，进入已加入课程和教师下发任务；遇到知识点薄弱或想主动学习时进入“自主学习”；科研项目实践是拓展入口，项目申报演示中可作为能力延展说明。",
    )
    add_table(
        doc,
        ["入口", "什么时候点", "进入后主要做什么"],
        [
            ["我的课程", "教师已发布课程或作业，需要完成课程任务时", "进入课程工作台，查看课程任务、收藏夹、学习画像和知识图谱"],
            ["自主学习", "想按知识点补课，或从画像推荐继续学习时", "生成讲解、练习、笔记、知识卡片、思维导图或 PPT 大纲"],
            ["科研项目实践", "需要展示平台扩展到项目实践和研究训练时", "查看项目建议、资料、阶段任务和研究辅助能力"],
            ["新手引导", "第一次登录、忘记主线或需要给评审演示时", "按 6 步提示理解学习闭环、AI 诊断、画像、引用和资料沉淀"],
        ],
        [1.25, 2.25, 3.15],
    )
    add_picture(doc, "student-01-entry-clean.png", "图 2 学生端服务入口")
    add_picture(doc, "student-01-entry.png", "图 3 学生端内置新手引导")

    add_heading(doc, "三 学生完成一次课程任务", 1)
    add_para(doc, "推荐给新用户的第一条路线是“我的课程 -> 课程工作台 -> 课程任务 -> 任务工作台 -> AI 诊断与提示 -> 保存总结”。这条线最能体现平台不是普通判题系统，而是助学闭环。")
    add_numbered(
        doc,
        [
            "在学生入口点击“进入课程”，进入课程工作台。",
            "先看左侧课程菜单和页面中的最近任务，确认当前课程、授课教师和未完成任务。",
            "进入“课程任务”，按状态筛选未完成、进行中或待修正任务。",
            "打开任务工作台，在中间代码编辑器完成代码，使用“运行代码”做自检，再点击“提交判题”。",
            "右侧先看系统验证结果，再看 AI 学习助手的诊断、引用来源、置信度和分层提示。",
            "根据提示修改并重新提交。通过后查看学习总结，保存到我的资料，并让学习画像更新。",
        ],
    )
    add_picture(doc, "student-02-course-home.png", "图 4 学生课程工作台")
    add_picture(doc, "student-03-course-tasks.png", "图 5 学生课程任务列表")
    add_picture(doc, "student-04-task-workspace.png", "图 6 学生任务工作台与 AI 学习助手")

    add_heading(doc, "四 学生自主学习和资料沉淀", 1)
    add_para(doc, "当学生不想从教师任务进入，或者任务卡住需要补知识点时，使用“自主学习”。这个入口把学习画像、知识库、AI 助学和资源中心放在一条线上，适合快速生成复习材料并继续练习。")
    add_table(
        doc,
        ["场景", "建议入口", "产出"],
        [
            ["不知道今天补什么", "自主学习首页的每日学习建议", "推荐主题、外部资源、练习或讲解"],
            ["概念没懂", "AI 助学", "分步解释、引用来源、下一步动作"],
            ["需要复习资料", "资源中心 / 我的资料", "笔记、知识卡片、错题总结、思维导图、PPT 大纲"],
            ["想看自己为什么被推荐", "个人画像", "薄弱知识点、高频错因、提示依赖和证据来源"],
        ],
        [1.55, 1.75, 2.95],
    )
    add_picture(doc, "student-05-self-study.png", "图 7 自主学习首页")
    add_picture(doc, "student-06-ai-tutor.png", "图 8 AI 助学页面")
    add_picture(doc, "student-07-library.png", "图 9 我的资料与资源中心")

    add_heading(doc, "五 教师端入口地图", 1)
    add_para(doc, "教师端的主线是“我的课程 -> 进入课程 -> 建班和邀请学生 -> 管理资料与知识库 -> 发布任务 -> 查看提交和批改 -> 学情分析”。第一次使用时不要从所有菜单同时展开，先围绕一门课程走完整闭环。")
    add_table(
        doc,
        ["教师入口", "主要用途", "新用户优先级"],
        [
            ["工作台首页", "查看课程、待处理任务和通知", "先看"],
            ["我的课程", "进入或管理某一门课程", "先看"],
            ["课程工作空间", "查看课程概览、班级、任务和知识库状态", "先看"],
            ["班级 / 邀请", "绑定授课班级并生成学生加入入口", "建课后用"],
            ["任务管理", "创建、编辑、发布作业或测验，预览学生视角", "高频使用"],
            ["成绩 / 批改", "查看规则分、AI 建议、教师反馈和发布状态", "有提交后用"],
            ["学情分析", "查看知识点掌握、错因统计和教学干预", "批改后用"],
        ],
        [1.5, 3.1, 1.4],
    )
    add_picture(doc, "teacher-02-courses.png", "图 10 教师我的课程")
    add_picture(doc, "teacher-03-course-workspace.png", "图 11 教师课程工作空间")

    add_heading(doc, "六 教师发布任务与查看学情", 1)
    add_para(doc, "教师日常最常用的是任务管理和学情分析。任务管理负责把作业、测验和编程题下发到班级；学生提交后，平台把沙箱结果、AI 诊断、提示使用和成绩反馈回传到教师端，形成班级层面的学情判断。")
    add_numbered(
        doc,
        [
            "进入“我的课程”，选择本学期正在进行的课程。",
            "在课程工作空间中进入“任务管理”，新建或编辑作业，设置题型、知识点、开始时间、截止时间和下发班级。",
            "发布前使用“预览学生视角”，确认学生看到的标题、说明、知识点和提示策略。",
            "学生提交后进入“查看成绩”或“学生成绩”，处理 AI 预评、教师反馈和最终发布。",
            "进入“学情分析”，按班级、任务或知识点查看完成率、平均分、逾期率、平均提示等级、风险学生和薄弱知识点。",
            "根据学情分析下发专项练习、推送复习提醒或发起课堂讨论。",
        ],
    )
    add_picture(doc, "teacher-04-tasks.png", "图 12 教师任务管理")
    add_picture(doc, "teacher-06-analytics.png", "图 13 教师学情分析")

    add_heading(doc, "七 教师建设课程知识来源", 1)
    add_para(doc, "教师端的资料管理和课程知识图谱决定 AI 回答是否有可追溯来源。项目演示时可以强调：学生端看到的 AI 诊断、AI 助学回答和自主学习资料，不应脱离教师确认过的课程资料和知识点结构。")
    add_bullets(
        doc,
        [
            "资料管理用于上传讲义、课件、实验指导和链接，并控制是否对学生可见。",
            "课程知识库用于把资料变成 AI 可检索的知识来源，便于回答时展示引用。",
            "课程知识图谱用于维护知识点、前置关系和资料证据，支撑学习路径推荐和薄弱点定位。",
            "AI 助教页面适合测试课程知识库是否能回答学生常见问题。",
        ],
    )
    add_picture(doc, "teacher-07-materials.png", "图 14 教师资料管理")
    add_picture(doc, "teacher-08-graph.png", "图 15 课程知识图谱")

    add_heading(doc, "八 推荐演示路线", 1)
    add_para(doc, "如果需要给评审或新用户做 5 到 8 分钟演示，建议不要逐个菜单介绍，而是按下列路线讲一个完整闭环。")
    add_table(
        doc,
        ["顺序", "演示动作", "要强调的能力"],
        [
            ["1", "学生登录，查看服务入口和新手引导", "平台把入口收敛为课程、自主学习和项目实践"],
            ["2", "学生进入数据结构课程工作台并打开任务", "任务来自教师发布，关联人工智能专业支撑课程"],
            ["3", "学生在任务工作台运行或提交代码", "先呈现沙箱与测试事实，再给 AI 诊断"],
            ["4", "学生查看分层提示和引用来源", "AI 不直接替代答案，回答有来源和置信度"],
            ["5", "学生进入自主学习，生成讲解或资料", "自主学习与画像、知识库、资料中心联动"],
            ["6", "教师进入课程任务管理", "教师能发布任务并预览学生视角"],
            ["7", "教师查看学情分析", "学生行为回传为知识点掌握、错因和干预建议"],
            ["8", "教师展示资料管理或知识图谱", "垂类 AI 由课程资料和知识结构支撑"],
        ],
        [0.7, 2.45, 3.05],
    )

    add_heading(doc, "九 常见迷路点", 1)
    add_table(
        doc,
        ["问题", "快速判断", "处理方式"],
        [
            ["登录后不知道去哪", "看右上角角色和首页标题", "学生优先点“我的课程”；教师优先点“我的课程”或“进入工作台”"],
            ["学生看不到课程任务", "是否进入了具体课程工作台", "从学生入口点“我的课程”，再进入某门课程的“课程任务”"],
            ["AI 助学和自主学习分不清", "AI 助学负责问答，自主学习负责按知识点组织学习流程", "概念问题点 AI 助学；想生成练习、笔记和学习计划点自主学习"],
            ["教师菜单很多", "先围绕一门课程处理", "按“课程 -> 班级 -> 资料 -> 任务 -> 批改 -> 学情”的顺序走"],
            ["批改页为空", "当前任务可能暂无学生提交", "先用任务管理查看是否有提交，再进入成绩或学情页面"],
            ["项目申报时怕讲散", "不要按菜单讲", "按第八节演示路线讲学生闭环和教师支撑闭环"],
        ],
        [1.5, 2.1, 2.55],
    )

    add_para(
        doc,
        "这份引导材料的核心使用方式是：先让学生和教师各自找到唯一主线，再把 AI 诊断、引用来源、资料沉淀和画像更新放回这条主线中理解。这样即使平台入口较多，新用户也能先完成一个可验证的学习闭环。",
    )

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)


if __name__ == "__main__":
    build()
    print(OUT_PATH)
